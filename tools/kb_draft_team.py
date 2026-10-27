"""把组员资料（2026-10-27）中已逐条核实的内容，整理成知识库初稿 Excel。

输出：docs/知识库初稿_组员资料核实版.xlsx（以 docs/知识库填写范本.xlsx 为底）。
所有条目「未审核」——须由医学背景人员审核后才可改为「已审核」。
出处已用 PubMed/Europe PMC 核对作者、期刊、年份（组员原稿有题录错误，此处为更正后版本）。
用法：backend\\.venv\\Scripts\\python tools\\kb_draft_team.py
"""
from __future__ import annotations

import shutil
import sys
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = ROOT / "docs" / "知识库填写范本.xlsx"
OUT = ROOT / "docs" / "知识库初稿_组员资料核实版.xlsx"

ADA24 = ("ADA. 11. Chronic Kidney Disease and Risk Management: Standards of Care in Diabetes—2024. "
         "Diabetes Care 2024;47(Suppl 1)", "2024")
ADA24_URL = "https://pmc.ncbi.nlm.nih.gov/articles/PMC10725805/"

# 编号, 指标, 标题, 内容, 关键词, 来源, 年份, 段落, 链接, 授权, 备注
ENTRIES = [
    ("kb-ket-101", "KET", "尿酮阴性为什么不能完全放心",
     "尿酮试纸主要测的是乙酰乙酸，而血液里含量最多的酮体是β-羟丁酸。酮症早期尿酮可能还没变化，治疗恢复期尿酮又可能比血酮晚消退。"
     "所以尿酮阴性不能完全排除酮症；如果同时有恶心、呕吐、腹痛、呼吸又深又快、精神差等不适，应及时就医，由医生决定是否查血酮。",
     "尿酮阴性,血酮,酮症,呕吐,呼吸急促",
     "Huang J, et al. Update on Measuring Ketones. J Diabetes Sci Technol 2024;18(5)", "2024", "摘要/正文",
     "https://pmc.ncbi.nlm.nih.gov/articles/PMC11089855/", "仅摘要引用",
     "辅证：Kuru B, et al. Turk J Emerg Med 2014（PMID 27331169）急诊高血糖患者中，部分尿酮阴性者指尖血酮为阳性"),
    ("kb-ket-102", "KET", "尿酮阳性代表什么",
     "尿酮阳性表示身体正在分解脂肪来供能，常见于胰岛素不足、长时间没吃东西、呕吐、饮酒等情况。"
     "急诊研究发现尿酮试纸很少漏掉糖尿病酮症酸中毒，但阳性的人里真正是酮症酸中毒的只占一部分，"
     "所以阳性需要医生结合血酮、血气等检查判断，不能只凭试纸下结论。尿糖与尿酮同时阳性时应尽快就医。",
     "尿酮阳性,酮体,不吃饭,喝酒,酮症酸中毒",
     "Arora S, et al. Diagnostic accuracy of point-of-care testing for DKA at emergency-department triage. Diabetes Care 2011;34(4)",
     "2011", "摘要", "https://pubmed.ncbi.nlm.nih.gov/21307381/", "仅摘要引用",
     "辅证：Schwab TM, et al. Ann Emerg Med 1999（PMID 10459090）"),
    ("kb-glu-101", "GLU", "尿糖和血糖为什么对不上",
     "血糖超过肾脏能回收的程度时，糖才会漏到尿里，而每个人肾脏开始“漏糖”的血糖水平不一样，同样的血糖下排出的糖量可以相差好几倍。"
     "因此尿糖不能换算成血糖值；尿糖阴性也不代表血糖一定正常。血糖控制得好不好，要看血糖或糖化血红蛋白。",
     "尿糖,血糖,肾糖阈,换算",
     "Wolf S, et al. Renal glucose excretion and tubular reabsorption rate related to blood glucose in subjects with type 2 diabetes. Horm Metab Res 2009;41(8):600-604",
     "2009", "摘要", "https://pubmed.ncbi.nlm.nih.gov/19418417/", "仅摘要引用",
     "组员原稿把第一作者误写为 Fuchs W，已更正"),
    ("kb-glu-102", "GLU,BLD,NIT", "维生素C会影响尿检结果",
     "大量维生素C（例如维生素C补充剂、泡腾片）可能让试纸上的尿糖、隐血、亚硝酸盐出现假阴性，也就是本来有却测不出来；"
     "不同品牌试纸受影响的程度不同。如果检查前几天吃过维生素C，复诊时可以告诉医生。",
     "维生素C,维C,泡腾片,假阴性,补充剂",
     "Unic A, et al. Ascorbic acid—A black hole of urine chemistry screening. J Clin Lab Anal 2018;32(5)",
     "2018", "摘要", "https://pubmed.ncbi.nlm.nih.gov/29479748/", "仅摘要引用", ""),
    ("kb-glu-103", "GLU", "SGLT2 抑制剂为什么会让尿糖阳性",
     "SGLT2 抑制剂（如达格列净、恩格列净、卡格列净）的作用就是让肾脏少回收糖、多从尿里排糖，所以即使血糖不高，尿糖也常是阳性。"
     "这是药物本身的作用，不代表控糖失败。用药期间如果出现恶心、呕吐、腹痛、呼吸深快等不适，应及时就医。",
     "达格列净,恩格列净,卡格列净,SGLT2,降糖药,尿糖阳性",
     "Gallo LA, Wright EM, Vallon V. Probing SGLT2 as a therapeutic target for diabetes: basic physiology and consequences. Diab Vasc Dis Res 2015;12(2)",
     "2015", "摘要", "https://pubmed.ncbi.nlm.nih.gov/25616707/", "仅摘要引用",
     "组员原稿期刊误写为 Diabetes Metab Syndr Obes，已更正；用药期不适提示另见 kb-ket-101 出处"),
    ("kb-uacr-101", "UACR", "UACR 升高一次，为什么要复查",
     "尿白蛋白每次检查之间的波动很大，一次升高不能说明一直偏高。指南建议在 3～6 个月内复查，"
     "3 次中有 2 次异常，才认为白蛋白尿是持续存在的。所以首次升高时，最重要的是按医生安排复查。",
     "复查,UACR升高,尿微量白蛋白,确认",
     ADA24[0], ADA24[1], "11 章 筛查部分", ADA24_URL, "仅摘要引用", ""),
    ("kb-uacr-102", "UACR,LEU,NIT", "哪些情况会让 UACR 暂时升高",
     "指南列出：检查前 24 小时内运动、感染、发烧、月经期、血糖明显偏高、血压明显偏高、心力衰竭，都可能让 UACR 暂时升高，并不一定是肾脏受损。"
     "留尿前避免剧烈运动；正在发烧、感染或月经期时，可以和医生商量是否延后检查。",
     "运动,发烧,感染,月经,经期,血压高,暂时升高",
     ADA24[0], ADA24[1], "11 章 筛查部分", ADA24_URL, "仅摘要引用", ""),
    ("kb-uacr-103", "UACR", "运动后多久留尿比较好",
     "一项糖尿病患者研究发现，步行约 1 公里后 1 小时，部分人的 UACR 会升高；到 2 小时时，已和早晨第一次尿的结果没有明显差别。"
     "因此留尿前最好避免运动，或运动后间隔一段时间再留尿。",
     "运动,走路,留尿,晨尿",
     "Koh KH, et al. Effect of exercise on albuminuria in people with diabetes. Nephrology (Carlton) 2011;16(8):704-709",
     "2011", "摘要", "https://pubmed.ncbi.nlm.nih.gov/21854502/", "仅摘要引用", "样本量小（糖尿病 35 例）"),
    ("kb-uacr-104", "UACR", "UACR 多久查一次",
     "指南建议：2 型糖尿病患者从确诊开始、1 型糖尿病患者病程满 5 年起，至少每年检查一次 UACR，并同时评估肾功能（eGFR）。"
     "如果已经有异常，复查频率由医生根据情况决定。",
     "多久查,每年,筛查,肾功能,eGFR",
     ADA24[0], ADA24[1], "推荐 11.1a", ADA24_URL, "仅摘要引用", ""),
    ("kb-pro-101", "PRO,UACR", "尿蛋白和 UACR 有什么不同",
     "尿常规里的尿蛋白是试纸半定量结果，对少量白蛋白不够敏感；UACR 是实验室定量检查，能发现更早期的变化。"
     "所以尿蛋白阴性不等于 UACR 正常，糖尿病的肾脏筛查以 UACR 和肾功能（eGFR）为准。",
     "尿蛋白,UACR,区别,阴性",
     "Mejia JR, et al. Diagnostic accuracy of urine dipstick testing for albumin-to-creatinine ratio and albuminuria: a systematic review and meta-analysis. Heliyon 2021;7(11)",
     "2021", "摘要", "https://pubmed.ncbi.nlm.nih.gov/34765776/", "仅摘要引用",
     "该荟萃分析自评证据确定性很低；筛查以 UACR 为准的依据见 ADA 2024 第 11 章"),
    ("kb-leu-101", "LEU,NIT", "亚硝酸盐阴性就没有感染吗",
     "不一定。糖尿病患者的研究显示：亚硝酸盐阳性时比较可靠，但很多感染者亚硝酸盐是阴性；白细胞比较敏感，但也会漏掉一部分。"
     "是否有尿路感染，需要医生结合尿频、尿急、尿痛等症状和尿培养来判断。",
     "亚硝酸盐,白细胞,尿路感染,尿痛,尿频",
     "Mohanna AT, et al. The Sensitivity and Specificity of White Blood Cells and Nitrite in Dipstick Urinalysis in Diabetic and Non-Diabetic Patients. Cureus 2021;13(6)",
     "2021", "摘要", "https://pubmed.ncbi.nlm.nih.gov/34249577/", "仅摘要引用", "单中心横断面研究"),
    ("kb-bld-101", "BLD", "糖尿病患者尿里有血怎么理解",
     "典型的糖尿病肾病比较少出现明显的血尿。尿隐血阳性可能来自月经、运动、泌尿系统等多种原因，"
     "医生可能会安排复查或进一步检查来找原因。复诊时可以问医生：这次隐血需不需要复查、要不要做其他检查。",
     "隐血,血尿,尿里有血,红细胞",
     "O'Neill WM, et al. Hematuria and red cell casts in typical diabetic nephropathy. Am J Med 1983;74(3)",
     "1983", "摘要", "https://pubmed.ncbi.nlm.nih.gov/6829589/", "仅摘要引用",
     "文献较旧、样本小；辅证 Lin HY, et al. Am J Med Sci 2018（PMID 30286822）"),
    ("kb-gen-101", "通用", "尿液放久了结果会不准",
     "尿液在室温下放太久，酸碱度、尿糖、尿蛋白等试纸结果会逐渐变得不可靠。留尿后应尽快送检；"
     "如果报告结果和自己的情况差别很大，可以问医生是否需要用新鲜尿复查。",
     "放久,送检,留尿,不准",
     "Bohn MK, et al. Evaluation of pre-analytical factors impacting urine test strip and chemistry results. Clin Chem Lab Med 2025;63(6):1169-1178",
     "2025", "摘要", "https://pubmed.ncbi.nlm.nih.gov/39788127/", "仅摘要引用", "研究每组样本 10–20 份"),
    ("kb-ph-101", "PH", "尿酸碱度（pH）受什么影响",
     "尿 pH 会受饮食影响，例如肉类等蛋白质吃得多时尿液会偏酸。一次 pH 偏低或偏高，不能用来判断酮症或酸中毒，"
     "这类情况要靠血液检查。",
     "酸碱度,pH,偏酸,饮食,吃肉",
     "Miki A, et al. Urinary pH reflects dietary acid load in patients with type 2 diabetes. J Clin Biochem Nutr 2017;61(1)",
     "2017", "摘要", "https://pubmed.ncbi.nlm.nih.gov/28751813/", "仅摘要引用", "2 型糖尿病 173 例横断面研究"),
]


def main() -> int:
    sys.stdout.reconfigure(encoding="utf-8")
    shutil.copy(TEMPLATE, OUT)
    wb = load_workbook(OUT)
    ws = wb["知识条目"]
    for r, (eid, codes, title, text, kw, src, ver, sec, url, lic, note) in enumerate(ENTRIES, start=4):
        row = [eid, codes, title, text, kw, src, ver, sec, url, lic, "未审核", "", "开发（据组员资料）", note]
        for c, v in enumerate(row, start=1):
            ws.cell(row=r, column=c, value=v)
    wb.save(OUT)
    print(f"已生成 {OUT.relative_to(ROOT)}（{len(ENTRIES)} 条，全部未审核）")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
