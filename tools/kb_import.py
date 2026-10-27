"""把组员填写的知识库 Excel 导入为正式知识库 backend/app/knowledge_base/kb.yaml。
会检查必填栏位、指标代码、编号重复、审核与授权状态；有错误时不写入，并逐条列出（第几行、哪一栏）。

用法：
  backend\\.venv\\Scripts\\python tools\\kb_import.py 知识库.xlsx            # 检查并导入
  backend\\.venv\\Scripts\\python tools\\kb_import.py 知识库.xlsx --check    # 只检查，不写入
导入后重新启动后端即生效。
"""
import argparse
import re
import sys
from datetime import datetime
from pathlib import Path

import yaml
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "backend" / "app" / "knowledge_base" / "kb.yaml"
VALID_CODES = {"GLU", "KET", "PRO", "LEU", "NIT", "BLD", "SG", "PH", "UACR"}
LICENSES = {"可公开使用", "已获授权", "仅摘要引用", "待确认"}
REVIEWS = {"未审核", "已审核"}
REQUIRED = ["编号", "适用指标", "标题", "内容", "来源名称", "版本或年份", "授权情况", "审核状态"]
FORBIDDEN = re.compile(r"(建议|可以|应该)(您|你)?(停药|减药|换药|加药|调整剂量)|(您|你)(已经)?(确诊|患有|得了)")


def split(v) -> list[str]:
    return [x.strip() for x in re.split(r"[,，、;；\s]+", str(v or "")) if x.strip()]


def main() -> int:
    sys.stdout.reconfigure(encoding="utf-8")
    ap = argparse.ArgumentParser()
    ap.add_argument("xlsx")
    ap.add_argument("--check", action="store_true", help="只检查，不写入")
    ap.add_argument("--out", default=str(OUT), help="输出位置（测试用）")
    args = ap.parse_args()

    ws = load_workbook(args.xlsx, data_only=True)["知识条目"]
    header = [str(c.value).strip() if c.value else "" for c in ws[1]]
    missing_cols = [c for c in REQUIRED if c not in header]
    if missing_cols:
        print(f"✗ 缺少栏位：{'、'.join(missing_cols)}（请使用 docs/知识库填写范本.xlsx）")
        return 1
    col = {name: i for i, name in enumerate(header)}

    errors, warnings, entries, seen = [], [], [], set()
    skipped = 0
    for r, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        get = lambda k: (str(row[col[k]]).strip() if k in col and row[col[k]] is not None else "")  # noqa: E731
        if not any(v not in (None, "") for v in row):
            continue
        if get("编号").startswith("示例"):
            skipped += 1
            continue
        for k in REQUIRED:
            if not get(k):
                errors.append(f"第 {r} 行「{k}」未填写")
        eid = get("编号")
        if eid in seen:
            errors.append(f"第 {r} 行编号「{eid}」重复")
        seen.add(eid)
        codes = split(get("适用指标"))
        bad = [c for c in codes if c.upper() not in VALID_CODES and c != "通用"]
        if bad:
            errors.append(f"第 {r} 行「适用指标」有无效代码：{'、'.join(bad)}（可用：{'、'.join(sorted(VALID_CODES))}、通用）")
        if get("授权情况") and get("授权情况") not in LICENSES:
            errors.append(f"第 {r} 行「授权情况」应为：{'、'.join(LICENSES)}")
        if get("审核状态") and get("审核状态") not in REVIEWS:
            errors.append(f"第 {r} 行「审核状态」应为：未审核 或 已审核")
        if get("审核状态") == "已审核" and not get("审核人"):
            errors.append(f"第 {r} 行已审核但未填「审核人」")
        content = get("内容")
        if FORBIDDEN.search(content):
            errors.append(f"第 {r} 行「内容」含诊断或用药调整表述，请改写")
        if len(content) > 400:
            warnings.append(f"第 {r} 行「内容」较长（{len(content)} 字），建议拆成多条")
        if get("授权情况") == "待确认":
            warnings.append(f"第 {r} 行授权待确认，正式发布前须确认")
        if not get("页码或段落"):
            warnings.append(f"第 {r} 行未填「页码或段落」，可追溯性较弱")
        source = get("来源名称") + (f"（{get('页码或段落')}）" if get("页码或段落") else "")
        entries.append({
            "id": eid,
            "title": get("标题"),
            "codes": [c.upper() for c in codes if c != "通用"],
            "keywords": split(get("关键词")),
            "source": source,
            "source_version": get("版本或年份"),
            "source_url": get("原文链接"),
            "authorization": get("授权情况"),
            "review_status": get("审核状态"),
            "reviewer": get("审核人"),
            "author": get("填写人"),
            "text": content,
        })

    print(f"读取 {len(entries)} 条（跳过示例 {skipped} 条）")
    for w in warnings:
        print("  ⚠", w)
    if errors:
        print(f"\n✗ 发现 {len(errors)} 个错误，未导入：")
        for e in errors:
            print("  ✗", e)
        return 1
    if not entries:
        print("✗ 没有可导入的条目（示例行会被跳过）")
        return 1
    if args.check:
        print("✓ 检查通过（--check 模式，未写入）")
        return 0
    data = {
        "version": f"kb-{datetime.now():%Y%m%d-%H%M}",
        "review_status": "mixed" if any(e["review_status"] == "未审核" for e in entries) else "reviewed",
        "source_file": Path(args.xlsx).name,
        "entries": entries,
    }
    out = Path(args.out)
    out.write_text("# 正式知识库（由 tools/kb_import.py 从 Excel 导入，请勿手动编辑）\n"
                   + yaml.safe_dump(data, allow_unicode=True, sort_keys=False), encoding="utf-8")
    reviewed = sum(e["review_status"] == "已审核" for e in entries)
    print(f"✓ 已导入 {len(entries)} 条到 {out}（已审核 {reviewed} 条）。重新启动后端后生效。")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
