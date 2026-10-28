"""以竞赛提供的《AIC·AI+医学技术报告参考大纲.docx》为模板生成技术方案（保留封面、目录与标题样式）。
用法（需 python-docx）：python tools/aic_report/build_report.py <模板.docx> <输出.docx>
之后用 Word 更新目录并导出 PDF（见 export_pdf.ps1）。
"""
from __future__ import annotations

import sys
from pathlib import Path

import docx
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

sys.path.insert(0, str(Path(__file__).parent))
import content as C  # noqa: E402

FIG = Path(__file__).parent / "figures"
BODY_PT, SMALL_PT = 12, 10.5  # 正文小四号；表格与图题五号


def set_font(run, size=BODY_PT, bold=None, east="宋体", west="Times New Roman"):
    run.font.size = Pt(size)
    run.font.name = west
    run.font.color.rgb = RGBColor(0, 0, 0)  # 全文黑色
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    fonts.set(qn("w:eastAsia"), east)
    fonts.set(qn("w:ascii"), west)
    fonts.set(qn("w:hAnsi"), west)
    if bold is not None:
        run.bold = bold


def fmt(p, indent=True, align=None, before=0, after=0):
    pf = p.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.first_line_indent = Pt(BODY_PT * 2) if indent else Pt(0)
    # 模板的正文样式以「字符」为单位缩进（firstLineChars），会覆盖上面的设定：一并写明
    ind = p._p.get_or_add_pPr().get_or_add_ind()
    ind.set(qn("w:firstLineChars"), "200" if indent else "0")
    if align is not None:
        p.alignment = align


def body(doc, text, **kw):
    p = doc.add_paragraph(style="Normal")
    fmt(p, **kw)
    set_font(p.add_run(text))
    return p


def caption(doc, text, before=3, after=6):
    p = doc.add_paragraph(style="Normal")
    fmt(p, indent=False, align=WD_ALIGN_PARAGRAPH.CENTER, before=before, after=after)
    set_font(p.add_run(text), size=SMALL_PT, bold=True)
    return p


def set_cell(cell, text, bold=False, shade=None):
    cell.text = ""
    p = cell.paragraphs[0]
    fmt(p, indent=False, align=WD_ALIGN_PARAGRAPH.CENTER if bold else WD_ALIGN_PARAGRAPH.LEFT)
    set_font(p.add_run(text), size=SMALL_PT, bold=bold)
    if shade:
        tcpr = cell._element.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), shade)
        tcpr.append(shd)


def table(doc, title, head, rows, widths):
    caption(doc, title, before=6, after=3)
    t = doc.add_table(rows=1 + len(rows), cols=len(head))
    t.style = doc.styles["Table Grid"]
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for r, values in enumerate([head] + rows):
        for c, v in enumerate(values):
            cell = t.rows[r].cells[c]
            cell.width = Cm(widths[c])
            set_cell(cell, v, bold=(r == 0))  # 黑白：表头只加粗，不填底色
    # 表头在跨页时重复
    trpr = t.rows[0]._tr.get_or_add_trPr()
    th = OxmlElement("w:tblHeader")
    th.set(qn("w:val"), "true")
    trpr.append(th)
    doc.add_paragraph(style="Normal").paragraph_format.space_after = Pt(0)


def figure(doc, name, title, width):
    p = doc.add_paragraph(style="Normal")
    fmt(p, indent=False, align=WD_ALIGN_PARAGRAPH.CENTER, before=6)
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(str(FIG / name), width=Cm(width))
    caption(doc, title)


HEAD_PT = {1: 16, 2: 14, 3: 12}  # 一级三号、二级四号、三级小四号，均加粗


def heading(doc, text, level):
    p = doc.add_paragraph(style=f"Heading {level}")
    fmt(p, indent=False, before=12 if level == 1 else 6, after=6 if level == 1 else 3)
    p.paragraph_format.keep_with_next = True
    set_font(p.add_run(text), size=HEAD_PT[level], bold=True, east="宋体")
    return p


def remove(p):
    p._element.getparent().remove(p._element)


def build(template: str, out: str) -> None:
    doc = docx.Document(template)
    P = doc.paragraphs

    # —— 封面：作品名称与日期（团队名称/编号由参赛团队填写）——
    P[15].runs[1].text = f" {C.TITLE} "
    P[15].runs[1].font.size = Pt(14)
    P[21].runs[0].text = "日期：2026年10月28日"
    # 封面文字带有大纲级别，会被目录收录：去掉
    for p in P[:22]:
        for el in p._p.xpath(".//w:outlineLvl"):
            el.getparent().remove(el)
    # 目录标题旁的模板提示框
    mc = "{http://schemas.openxmlformats.org/markup-compatibility/2006}AlternateContent"
    for el in list(P[22]._p.iter(mc)):
        el.getparent().remove(el)
    # 文档网格会把「单倍行距」撑大：改为无网格
    for s in doc.sections:
        for g in s._sectPr.xpath("./w:docGrid"):
            for a in ("type", "charSpace"):
                g.attrib.pop(qn(f"w:{a}"), None)

    # —— 删除模板中的提示框与示例内容（保留携带分节符的段落）——
    sect1_end, sect2_end = P[38], P[44]
    abstract_title = P[39]
    to_remove = [P[i] for i in (32, 33, 34, 40, 41, 42, 43)] + P[45:]
    for p in to_remove:
        remove(p)

    # —— 作品简介 ——
    ap = sect2_end.insert_paragraph_before(C.ABSTRACT, style="Normal")
    fmt(ap)
    for r in ap.runs:
        set_font(r)
    set_font(abstract_title.runs[0], size=16, bold=True)

    # —— 正文 ——
    for block in C.BODY:
        kind = block[0]
        if kind in ("h1", "h2", "h3"):
            heading(doc, block[1], int(kind[1]))
        elif kind == "p":
            body(doc, block[1])
        elif kind == "formula":
            body(doc, block[1], indent=False, align=WD_ALIGN_PARAGRAPH.CENTER, before=3, after=3)
        elif kind == "table":
            table(doc, *block[1:])
        elif kind == "fig":
            figure(doc, *block[1:])

    # —— 附录 ——
    heading(doc, "七、附录", 1)
    for block in C.APPENDIX:
        if block[0] == "h2":
            heading(doc, block[1], 2)
        elif block[0] == "refs":
            for i, ref in enumerate(C.REFERENCES, 1):
                p = body(doc, f"[{i}] {ref}", indent=False)
                p.paragraph_format.left_indent = Pt(BODY_PT * 2)
                p.paragraph_format.first_line_indent = Pt(-BODY_PT * 2)
                for r in p.runs:
                    set_font(r, size=SMALL_PT)
        else:
            body(doc, block[1])

    # 页面设置：按模板要求统一页边距
    for s in doc.sections:
        s.top_margin = s.bottom_margin = Cm(2.5)
        s.left_margin = s.right_margin = Cm(3.0)
        s.header_distance = s.footer_distance = Cm(1.5)
    doc.save(out)
    print("已生成", out)


if __name__ == "__main__":
    build(sys.argv[1], sys.argv[2])
