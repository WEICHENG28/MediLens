"""生成知识库 Excel 填写范本：docs/知识库填写范本.xlsx
用法：backend\\.venv\\Scripts\\python tools\\kb_template.py
"""
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "知识库填写范本.xlsx"

COLUMNS = [
    ("编号", 14, "必填。格式 kb-指标-序号，如 kb-uacr-001；不可重复"),
    ("适用指标", 16, "必填。填指标代码，多个用逗号分隔，如 UACR,LEU；不针对特定指标填「通用」"),
    ("标题", 24, "必填。如「UACR 是什么」「感染为什么会影响 UACR」"),
    ("内容", 60, "必填。用通俗中文改写（约 50–200 字），不要整段复制原文；不得写诊断或停药换药建议"),
    ("关键词", 22, "选填。患者可能用的说法，逗号分隔，如 尿微量白蛋白,肾功能"),
    ("来源名称", 30, "必填。指南/教材/网站名称，如《中国糖尿病肾脏病防治指南（2021年版）》"),
    ("版本或年份", 12, "必填。如 2021、2025版"),
    ("页码或段落", 14, "选填但建议填写，如 第3章 3.2节、p.45"),
    ("原文链接", 30, "选填。网址或 DOI"),
    ("授权情况", 14, "必填。下拉选择：可公开使用 / 已获授权 / 仅摘要引用 / 待确认"),
    ("审核状态", 12, "必填。下拉选择：未审核 / 已审核（须由医学背景人员审核）"),
    ("审核人", 12, "审核状态为「已审核」时必填"),
    ("填写人", 10, "选填"),
    ("备注", 24, "选填"),
]
CODES = [("GLU", "尿葡萄糖"), ("KET", "尿酮体"), ("PRO", "尿蛋白"), ("LEU", "白细胞"), ("NIT", "亚硝酸盐"),
         ("BLD", "隐血"), ("SG", "比重"), ("PH", "酸碱度"), ("UACR", "尿白蛋白/肌酐比"), ("通用", "不针对特定指标")]
EXAMPLES = [
    ["示例-001", "UACR", "（示例）UACR 是什么",
     "（示例，导入时会自动跳过）在这里用通俗中文说明这个指标是什么、为什么要检查。请按原意改写，不要整段复制，并在右侧写清出处。",
     "尿微量白蛋白,肾", "（示例）《中国糖尿病肾脏病防治指南（2021年版）》", "2021", "（示例）第X章", "", "仅摘要引用", "未审核", "", "", ""],
    ["示例-002", "LEU,NIT,UACR", "（示例）感染为什么会影响 UACR",
     "（示例，导入时会自动跳过）说明干扰因素：例如尿路感染、发热、剧烈运动可能让结果暂时升高，需由医生结合复查判断。",
     "尿路感染,发烧,运动", "（示例）资料名称", "（示例）2025", "", "", "待确认", "未审核", "", "", ""],
]


def main() -> None:
    sys.stdout.reconfigure(encoding="utf-8")
    wb = Workbook()
    ws = wb.active
    ws.title = "知识条目"
    head_fill = PatternFill("solid", fgColor="00897B")
    for col, (name, width, _) in enumerate(COLUMNS, start=1):
        c = ws.cell(row=1, column=col, value=name)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = head_fill
        c.alignment = Alignment(horizontal="center", vertical="center")
        ws.column_dimensions[c.column_letter].width = width
    for r, row in enumerate(EXAMPLES, start=2):
        for col, v in enumerate(row, start=1):
            c = ws.cell(row=r, column=col, value=v)
            c.font = Font(color="808080", italic=True)
            c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "B2"
    for col_name, options in [("授权情况", "可公开使用,已获授权,仅摘要引用,待确认"), ("审核状态", "未审核,已审核")]:
        idx = [n for n, _, _ in COLUMNS].index(col_name) + 1
        letter = ws.cell(row=1, column=idx).column_letter
        dv = DataValidation(type="list", formula1=f'"{options}"', allow_blank=True)
        dv.add(f"{letter}2:{letter}1000")
        ws.add_data_validation(dv)

    guide = wb.create_sheet("填写说明")
    guide.column_dimensions["A"].width = 16
    guide.column_dimensions["B"].width = 90
    rows = [("MediLens 知识库填写说明", ""), ("", ""),
            ("用途", "知识条目只用于「解释」（是什么、为什么异常、干扰因素、该问医生什么），不用于判定风险；判定由规则负责。"),
            ("每条一行", "一条知识只讲一件事，约 50–200 字，通俗易懂。"),
            ("不要写", "诊断结论（如「您得了肾病」）、停药/换药/调整剂量建议、未经核实的数值。"),
            ("来源", "必须写清来源名称、版本或年份；建议写页码或段落。没有可靠来源的内容不要填。"),
            ("版权", "不要整段复制原文，按原意改写并标出处；授权不明确时「授权情况」选「待确认」。"),
            ("审核", "由医学背景人员审核后才能改为「已审核」，并填写审核人。"),
            ("示例行", "灰色斜体的「示例-」开头行会被导入工具自动跳过，可保留或删除。"),
            ("导入", "填好后交给开发同学执行：backend\\.venv\\Scripts\\python tools\\kb_import.py 文件路径.xlsx"),
            ("", ""), ("指标代码", "名称")] + CODES
    for r, (a, b) in enumerate(rows, start=1):
        guide.cell(row=r, column=1, value=a).font = Font(bold=(r == 1 or a in ("指标代码",)))
        guide.cell(row=r, column=2, value=b).alignment = Alignment(wrap_text=True)
    guide.cell(row=1, column=1).font = Font(bold=True, size=14)
    col_guide = wb.create_sheet("栏位说明")
    col_guide.column_dimensions["A"].width = 14
    col_guide.column_dimensions["B"].width = 90
    for r, (name, _, desc) in enumerate(COLUMNS, start=1):
        col_guide.cell(row=r, column=1, value=name).font = Font(bold=True)
        col_guide.cell(row=r, column=2, value=desc)
    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    print(f"已生成范本：{OUT}")


if __name__ == "__main__":
    main()
